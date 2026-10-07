import openpyxl
import re
import json

file_path = 'Analitica_Planeacion_2.0.xlsm'

def extract_sheet_dependencies():
    print("Loading workbook...")
    wb = openpyxl.load_workbook(file_path, data_only=False)
    
    dependencies = {}
    
    # Regex to find sheet names in formulas (e.g. 'Sheet Name'!A1 or SheetName!A1)
    sheet_regex = re.compile(r"'([^']+)'!|([a-zA-Z0-9_]+)!")
    
    for sheet_name in wb.sheetnames:
        dependencies[sheet_name] = set()
        ws = wb[sheet_name]
        
        formula_count = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 1000)):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    matches = sheet_regex.findall(cell.value)
                    for match in matches:
                        # match is a tuple like ('Sheet Name', '') or ('', 'SheetName')
                        ref_sheet = match[0] if match[0] else match[1]
                        if ref_sheet in wb.sheetnames and ref_sheet != sheet_name:
                            dependencies[sheet_name].add(ref_sheet)
                    formula_count += 1
                    
    # Convert sets to lists for JSON serialization
    for k in dependencies:
        dependencies[k] = list(dependencies[k])
        
    with open('sheet_dependencies.json', 'w', encoding='utf-8') as f:
        json.dump(dependencies, f, indent=4)
    print("Dependency extraction complete.")

if __name__ == '__main__':
    extract_sheet_dependencies()

